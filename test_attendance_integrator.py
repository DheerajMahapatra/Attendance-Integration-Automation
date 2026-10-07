"""
Tests for attendance_integrator.py  (run:  python test_attendance_integrator.py)

Covers the cases required by the project specification:
  1 student below 75%                7 missing subject
  2 student exactly at 75%           8 missing student
  3 regular student above 75%        9 duplicate course record
  4 regular student already at 95%  10 blank attendance
  5 multiple subjects                11 overall (weighted) calculation
  6 multiple semesters               12 semester (weighted) calculation
 13 the 75 - 78% / 95 - 98% bands    14 roster built without a MASTER file
"""

from __future__ import annotations

import contextlib
import io
import math
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

import attendance_integrator as ai


@contextlib.contextmanager
def argv_is(args):
    """Run with a known sys.argv (main() reads the folder argument from it)."""
    old = sys.argv
    sys.argv = ["attendance_integrator.py", *args]
    try:
        yield
    finally:
        sys.argv = old


@contextlib.contextmanager
def chdir(path):
    """Scan folder discovery inside `path` instead of the real project folder."""
    old = Path.cwd()
    os.chdir(path)
    try:
        yield
    finally:
        os.chdir(old)


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def make_master(path: Path, students=None) -> None:
    """Small MASTER with semester sections and one subject section (course code)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Master"

    ws["A1"] = "College Attendance Master"
    ws.merge_cells("A1:V1")

    ws["E2"] = "Semester -01"
    ws.merge_cells("E2:J2")
    ws["Q2"] = "Semester -02"
    ws.merge_cells("Q2:V2")
    ws["K2"] = "Sub: Mathematics BTE1001"
    ws.merge_cells("K2:P2")

    info_headers = ["Sl. No", "Student Name", "Enrolment No.", ""]
    for i, h in enumerate(info_headers, start=1):
        ws.cell(3, i).value = h

    semester_fields = ["Regular/Lateral/Ex-Regular", "Total No. of Classess",
                       "Total No of Actual  Clases Attended (A)",
                       "Total No of Additional Classes(B)",
                       "Total No. of Classess Attended(A+B)", "Current Attendance %"]
    subject_fields = ["Total No of Allotted Classes", "Total No. of Classess",
                      "Total No of Classes Attended- Actual (A)",
                      "Total No of Classes Added (B)",
                      "Total No of Classess Attended (A+B)", "Current Attendance %"]
    for offset, text in enumerate(semester_fields):
        ws.cell(3, 5 + offset).value = text
        ws.cell(3, 17 + offset).value = text
    for offset, text in enumerate(subject_fields):
        ws.cell(3, 11 + offset).value = text

    students = students or [(1, "ALPHA STUDENT", "AJU/1001"),
                            (2, "BETA STUDENT", "AJU/1002")]
    for i, (sr, name, usn) in enumerate(students):
        r = 4 + i
        ws.cell(r, 1).value = sr
        ws.cell(r, 2).value = name
        ws.cell(r, 3).value = usn
    wb.save(path)


def make_attendance(path: Path) -> None:
    headers = ["USN NO", "STUDENT TYPE", "STUDENT_NAME", "SEMESTER", "SECTION",
               "COURSE_CODE1", "ATTENDANCE_FOR1", "TOTAL_CLASS1", "ATTENDED_CLASS1",
               "PERCENTAGE1",
               "COURSE_CODE2", "ATTENDANCE_FOR2", "TOTAL_CLASS2", "ATTENDED_CLASS2",
               "PERCENTAGE2"]
    rows = [
        # normal student: one subject below 75%, one unknown course
        ["AJU/1001", "REGULAR", "ALPHA STUDENT", "I", "C",
         "BTE1001", "THEORY", 60, 40, 67,
         "BTE9999", "THEORY", 40, 20, 50],
        # blank course code + missing TOTAL_CLASS (must not invent anything)
        ["AJU/1001", "REGULAR", "ALPHA STUDENT", "I", "C",
         "", "THEORY", 30, 20, 67,
         "BTE1001", "THEORY", None, 10, 0],
        # student that does not exist in the MASTER
        ["AJU/7777", "REGULAR", "GHOST STUDENT", "I", "C",
         "BTE1001", "THEORY", 60, 45, 75],
        # student in MASTER but TOTAL_CLASS is blank -> nothing is invented
        ["AJU/1002", "REGULAR", "BETA STUDENT", "I", "C",
         "BTE1001", "THEORY", None, 10, 0,
         None, None, None, None, None],
        # second semester for the overall calculation
        ["AJU/1001", "REGULAR", "ALPHA STUDENT", "II", "C",
         "BTE1001", "THEORY", 100, 90, 90, None, None, None, None, None],
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


# --------------------------------------------------------------------------- #
# 1 - 4 : the attendance business rules
# --------------------------------------------------------------------------- #
class AttendanceRulesTest(unittest.TestCase):
    def test_1_student_below_75_percent(self):
        # Total 60, Actual 40 (66.67%) -> pushed inside the 75 - 78% band
        b, final, pct = ai.calculate_required_classes(60, 40, "REGULAR")
        self.assertEqual(b, 6)
        self.assertEqual(final, 46)
        self.assertAlmostEqual(pct, 76.67, places=2)
        self.assertTrue(ai.LOW_BAND[0] <= pct <= ai.LOW_BAND[1])

    def test_2_student_exactly_75_percent(self):
        b, final, pct = ai.calculate_required_classes(60, 45, "LATERAL")
        self.assertEqual(b, 0)
        self.assertEqual(final, 45)
        self.assertAlmostEqual(pct, 75.0)

    def test_3_regular_student_above_75_percent(self):
        # 50/60 = 83.33% -> pushed into the 95 - 98% band (B = 8 -> 96.67%)
        b, final, pct = ai.calculate_required_classes(60, 50, "REGULAR")
        self.assertEqual(b, 8)
        self.assertEqual(final, 58)
        self.assertAlmostEqual(pct, 96.67, places=2)
        self.assertTrue(ai.HIGH_BAND[0] <= pct <= ai.HIGH_BAND[1])

    def test_4_regular_student_at_95_percent_is_untouched(self):
        # already inside the band -> nothing is added (never reduced either)
        b, final, pct = ai.calculate_required_classes(100, 95, "REGULAR")
        self.assertEqual(b, 0)
        self.assertEqual(final, 95)
        self.assertAlmostEqual(pct, 95.0)

    def test_4b_regular_student_96_point_67_is_untouched(self):
        b, final, pct = ai.calculate_required_classes(60, 58, "REGULAR")
        self.assertEqual(b, 0)
        self.assertEqual(final, 58)
        self.assertAlmostEqual(pct, 96.67, places=2)

    def test_lateral_student_gets_no_bonus(self):
        b, final, pct = ai.calculate_required_classes(60, 55, "LATERAL")
        self.assertEqual(b, 0)
        self.assertAlmostEqual(pct, 55 / 60 * 100)

    def test_ex_regular_student_gets_no_bonus(self):
        b, _final, pct = ai.calculate_required_classes(60, 55, "EX-REGULAR")
        self.assertEqual(b, 0)
        self.assertAlmostEqual(pct, 55 / 60 * 100)

    def test_10_blank_attendance_is_not_invented(self):
        self.assertIsNone(ai.calculate_subject_attendance(None, 10, "REGULAR"))
        self.assertIsNone(ai.calculate_subject_attendance(60, None, "REGULAR"))
        self.assertIsNone(ai.calculate_subject_attendance(0, 0, "REGULAR"))

    # ---------------------------------------------------------- the two bands
    @staticmethod
    def _reachable(total, actual, band) -> bool:
        """True when at least one class count lands inside the band."""
        low, high = band
        b_min = max(0, math.ceil(low / 100 * total - actual - 1e-9))
        b_max = math.floor(high / 100 * total - actual + 1e-9)
        return b_max >= b_min

    def test_below_75_lands_inside_75_78(self):
        low, high = ai.LOW_BAND
        for total in (30, 45, 60, 68, 100, 240, 365):
            for actual in range(0, math.ceil(total * 0.75)):
                res = ai.calculate_subject_attendance(total, actual, "REGULAR",
                                                      f"s{total}{actual}")
                self.assertGreaterEqual(res["percentage"], low - 0.01,
                                        f"{actual}/{total}")
                if self._reachable(total, actual, ai.LOW_BAND):
                    self.assertLessEqual(res["percentage"], high + 0.01,
                                         f"{actual}/{total}")

    def test_at_above_75_lands_inside_95_98(self):
        low, high = ai.HIGH_BAND
        for total in (30, 45, 60, 68, 100, 240, 365):
            for actual in range(math.ceil(total * 0.75), total + 1):
                res = ai.calculate_subject_attendance(total, actual, "REGULAR",
                                                      f"h{total}{actual}")
                current = actual / total * 100
                if current < low:
                    self.assertGreaterEqual(res["percentage"], low - 0.01,
                                            f"{actual}/{total}")
                    if self._reachable(total, actual, ai.HIGH_BAND):
                        self.assertLessEqual(res["percentage"], high + 0.01,
                                             f"{actual}/{total}")
                else:
                    self.assertEqual(res["added"], 0, f"{actual}/{total}")
                    self.assertAlmostEqual(res["percentage"], current, places=6)
                self.assertLessEqual(res["percentage"], 100.0)

    def test_never_above_100_and_never_negative(self):
        for total, actual in ((100, 100), (60, 60), (15, 11), (8, 6), (7, 0)):
            res = ai.calculate_subject_attendance(total, actual, "REGULAR", "x")
            self.assertLessEqual(res["final"], total)
            self.assertLessEqual(res["percentage"], 100.0)
            self.assertGreaterEqual(res["added"], 0)
            self.assertGreaterEqual(res["percentage"], 0.0)

    def test_spread_inside_a_band_is_deterministic(self):
        first = ai.calculate_subject_attendance(100, 80, "REGULAR",
                                                ai.spread_seed("AJU/1001", 1, "BTE1001"))
        again = ai.calculate_subject_attendance(100, 80, "REGULAR",
                                                ai.spread_seed("AJU/1001", 1, "BTE1001"))
        other = ai.calculate_subject_attendance(100, 80, "REGULAR",
                                                ai.spread_seed("AJU/1002", 1, "BTE1001"))
        self.assertEqual(first, again)                  # reproducible
        self.assertNotEqual(first["added"], other["added"])   # but not identical
        for res in (first, other):
            self.assertTrue(ai.HIGH_BAND[0] <= res["percentage"] <= ai.HIGH_BAND[1])

    def test_band_problem_reports_violations(self):
        self.assertIsNone(ai.band_problem(60, 40, 6, "REGULAR"))     # 76.67%
        self.assertIsNone(ai.band_problem(60, 45, 0, "LATERAL"))     # 75% kept
        self.assertIsNone(ai.band_problem(100, 95, 0, "REGULAR"))    # in band
        self.assertIsNotNone(ai.band_problem(60, 40, 0, "REGULAR"))  # still 66.67%
        self.assertIsNotNone(ai.band_problem(100, 90, 0, "REGULAR"))  # not raised
        self.assertIsNotNone(ai.band_problem(60, 45, 1, "LATERAL"))  # not allowed
        self.assertIsNotNone(ai.band_problem(100, 50, 100, "REGULAR"))  # > 100%


# --------------------------------------------------------------------------- #
# 5, 6, 11, 12 : weighted totals (never an average of percentages)
# --------------------------------------------------------------------------- #
class WeightedTotalsTest(unittest.TestCase):
    def test_5_multiple_subjects_are_weighted(self):
        # spec example: 45/60 and 20/40 -> actual 65 of 100 (never 62.5%)
        s1 = ai.calculate_subject_attendance(60, 45, "LATERAL")
        s2 = ai.calculate_subject_attendance(40, 20, "LATERAL")
        total = ai.calculate_semester_total([s1, s2])
        self.assertEqual(total["total"], 100)
        self.assertEqual(total["actual"], 65)          # summed, not averaged
        # 45/60 stays at 75% (LATERAL), 20/40 is pushed to 31/40 (77.5%)
        self.assertEqual(total["final"], 45 + s2["final"])
        self.assertTrue(ai.LOW_BAND[0] <= s2["percentage"] <= ai.LOW_BAND[1])
        # the percentage always comes from the class counts
        self.assertAlmostEqual(total["percentage"], total["final"] / total["total"] * 100)
        # raw attendance (before B) is 65/100 = 65%, not the 62.5% average
        self.assertAlmostEqual(total["actual"] / total["total"] * 100, 65.0)

    def test_12_semester_total_is_not_an_average(self):
        subjects = [ai.calculate_subject_attendance(100, 90, "LATERAL"),
                    ai.calculate_subject_attendance(20, 19, "LATERAL")]
        total = ai.calculate_semester_total(subjects)
        # 109 / 120 = 90.83%   -   average of 90% and 95% would be 92.5%
        self.assertAlmostEqual(total["percentage"], 109 / 120 * 100)
        self.assertNotAlmostEqual(total["percentage"], (90 + 95) / 2)

    def test_6_and_11_multiple_semesters_are_weighted(self):
        sem1 = ai.calculate_semester_total([ai.calculate_subject_attendance(156, 120, "LATERAL")])
        sem2 = ai.calculate_semester_total([ai.calculate_subject_attendance(180, 140, "LATERAL")])
        sem3 = ai.calculate_semester_total([ai.calculate_subject_attendance(170, 130, "LATERAL")])
        overall = ai.calculate_overall_attendance({1: sem1, 2: sem2, 3: sem3})
        self.assertEqual(overall["total"], 506)
        self.assertEqual(overall["final"], 390)
        self.assertAlmostEqual(overall["percentage"], 390 / 506 * 100)
        self.assertEqual(overall["semesters"], [1, 2, 3])

    def test_empty_semester_is_not_zero_percent(self):
        overall = ai.calculate_overall_attendance({1: None, 2: {"total": 0}})
        self.assertIsNone(overall)

        sem = ai.calculate_semester_total([])
        self.assertIsNone(sem)


# --------------------------------------------------------------------------- #
# normalisation + matching  (spec 2, 3, 4)
# --------------------------------------------------------------------------- #
class NormalisationTest(unittest.TestCase):
    def test_student_id_is_tolerant(self):
        self.assertEqual(ai.student_key("  AJU/230064 "), ai.student_key("aju/230064"))
        self.assertEqual(ai.student_key("AJU/230064"), ai.student_key("AJU / 230064"))
        self.assertNotEqual(ai.student_key("AJU/230064"), ai.student_key("AJU/230073"))

    def test_course_code_is_tolerant(self):
        self.assertEqual(ai.course_key(" bte21001 "), ai.course_key("BTE21001"))
        self.assertEqual(ai.normalize_course_code("bte 21001"), "BTE21001")

    def test_semester_from_many_formats(self):
        for value, expected in [("I", 1), ("II", 2), ("III", 3), ("IV", 4), ("V", 5),
                                ("VI", 6), ("VII", 7), ("1", 1), ("7", 7),
                                ("Semester -01", 1), ("SEM III", 3), ("2nd", 2),
                                ("Semester V", 5)]:
            self.assertEqual(ai.normalize_semester(value), expected, value)
        self.assertIsNone(ai.normalize_semester("ODD 2023-24"))
        self.assertIsNone(ai.normalize_semester(""))

    def test_semester_from_filename(self):
        self.assertEqual(ai.semester_from_filename("Semester_3.xlsx"), 3)
        self.assertEqual(ai.semester_from_filename("Attendance_Sem-II.xlsx"), 2)
        self.assertEqual(ai.semester_from_filename("Attendance_Sem-VII.xls"), 7)

    def test_student_type_detection(self):
        self.assertTrue(ai.is_regular("REGULAR"))
        self.assertFalse(ai.is_regular("LATERAL"))
        self.assertFalse(ai.is_regular("EX-REGULAR"))
        self.assertTrue(ai.is_regular(""))            # no type -> standard rule
        self.assertFalse(ai.is_regular("LATERAL ENTRY"))

    def test_band_examples_from_the_spec(self):
        # below 75% -> 75-78%, 75% and above -> 95-98%, never above 100
        for total, attended in ((60, 40), (100, 70), (240, 180), (100, 96)):
            res = ai.calculate_subject_attendance(total, attended, "REGULAR")
            current = attended / total * 100
            band = ai.LOW_BAND if current < ai.TARGET_PERCENT else ai.HIGH_BAND
            self.assertEqual(res["required_75"] > 0, current < ai.TARGET_PERCENT,
                             (total, attended))
            self.assertEqual(res["bonus_15"] > 0,
                             current >= ai.TARGET_PERCENT and current < band[0],
                             (total, attended))
            if current < ai.TARGET_PERCENT or current < band[0]:
                self.assertGreaterEqual(res["percentage"], band[0] - 0.01)
                self.assertLessEqual(res["percentage"], band[1] + 0.01)
            else:
                self.assertAlmostEqual(res["percentage"], current, places=6)
            self.assertLessEqual(res["percentage"], 100.0)

    def test_below_75_uses_required_classes_only(self):
        res = ai.calculate_subject_attendance(60, 40, "REGULAR")
        self.assertEqual(res["required_75"], 6)
        self.assertEqual(res["bonus_15"], 0)
        self.assertEqual(res["final"], 46)
        self.assertAlmostEqual(res["current"], 40 / 60 * 100, places=2)
        self.assertAlmostEqual(res["percentage"], 76.67, places=2)

    def test_actual_and_adjusted_are_separate(self):
        res = ai.calculate_subject_attendance(60, 40, "REGULAR")
        self.assertEqual(res["actual"], 40)            # untouched source value
        self.assertEqual(res["current"], 40 / 60 * 100)
        self.assertEqual(res["final"], 46)             # adjusted value
        self.assertEqual(res["added"], 6)


# --------------------------------------------------------------------------- #
# 7, 8, 9 : missing subject / missing student / duplicates
# --------------------------------------------------------------------------- #
class MatchingTest(unittest.TestCase):
    def _layout(self):
        ws = openpyxl.Workbook().active
        ws["A1"] = "x"
        ws["B3"] = "Student Name"
        ws["C3"] = "Enrolment No."
        ws["C4"] = "AJU/1001"
        return ai.parse_master_layout(ws)

    def test_8_missing_student(self):
        layout = self._layout()
        self.assertEqual(ai.match_student(layout, "AJU/1001"), 4)
        self.assertIsNone(ai.match_student(layout, "AJU/9999"))
        self.assertIsNone(ai.match_student(layout, ""))

    def test_7_missing_subject(self):
        layout = ai.MasterLayout(
            sheet_title="s", header_row=3, group_row=2, data_start=4,
            info_cols={"enrollment": 3, "name": 2, "type": None, "semester": None,
                       "father": None, "course": None, "batch": None, "section": None},
            semester_blocks={}, subject_blocks=[], subject_index={},
            student_rows={}, student_info={})
        block = ai.Block(start_col=5, end_col=10, label="Sub: Maths BTE1001",
                         kind="subject", semester=1, course_code="BTE1001",
                         course_key="BTE1001")
        layout.subject_blocks.append(block)
        layout.subject_index[(1, "BTE1001")] = block
        self.assertIsNotNone(ai.match_course(layout, 1, "BTE1001"))
        self.assertIsNone(ai.match_course(layout, 1, "BTE40404"))     # unknown course
        self.assertIsNone(ai.match_course(layout, 2, "BTE1001"))      # wrong semester
        self.assertTrue(layout.has_subject_catalog(1))
        self.assertFalse(layout.has_subject_catalog(2))

    def test_9_duplicate_records(self):
        def rec(total, attended):
            return ai.CourseRecord(
                source_file="f", source_sheet="s", row_number=1, slot=1,
                usn="AJU/1001", name="A", student_type="REGULAR", section="C",
                branch="", session="", scheme="", semester=1,
                course_code="BTE1001", course_key="BTE1001",
                attendance_for="THEORY", total=total, attended=attended,
                source_percentage=None)

        # identical duplicates -> keep one
        resolved = ai.resolve_duplicates([rec(60, 40), rec(60, 40)])
        kept = [r for r in resolved if r[1] == ai.KEEP]
        ignored = [r for r in resolved if r[1] == ai.IGNORE]
        self.assertEqual(len(kept), 1)
        self.assertEqual(len(ignored), 1)

        # duplicate without data -> keep the record that has data
        resolved = ai.resolve_duplicates([rec(None, None), rec(60, 40)])
        kept = [r for r in resolved if r[1] in (ai.KEEP, ai.CONFLICT_KEPT)]
        self.assertEqual(len(kept), 1)
        self.assertEqual(kept[0][0].total, 60)
        self.assertEqual(len([r for r in resolved if r[1] == ai.IGNORE]), 1)

        # different data -> keep the biggest total and flag the conflict
        resolved = ai.resolve_duplicates([rec(50, 30), rec(60, 40)])
        conflict = [r for r in resolved if r[1] == ai.CONFLICT_KEPT]
        self.assertEqual(len(conflict), 1)
        self.assertEqual(conflict[0][0].total, 60)
        self.assertIn("CONFLICT", conflict[0][2])

        # single record untouched
        resolved = ai.resolve_duplicates([rec(60, 40)])
        self.assertEqual(resolved[0][1], ai.KEEP)


# --------------------------------------------------------------------------- #
# end to end: build a tiny input folder, run the tool, check the workbook
# --------------------------------------------------------------------------- #
class EndToEndTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="attendance_test_"))
        (cls.tmp / "input").mkdir()
        (cls.tmp / "output").mkdir()
        make_master(cls.tmp / "input" / "MASTER.xlsx")
        make_attendance(cls.tmp / "input" / "Semester_1.xlsx")

        cls.old = (ai.INPUT_DIR, ai.OUTPUT_DIR, ai.OUTPUT_FILE)
        ai.INPUT_DIR = cls.tmp / "input"
        ai.OUTPUT_DIR = cls.tmp / "output"
        ai.OUTPUT_FILE = cls.tmp / "output" / "Integrated_Attendance.xlsx"
        with argv_is([]):                       # no folder argument -> use INPUT_DIR
            cls.exit_code = ai.main()
        cls.wb = openpyxl.load_workbook(ai.OUTPUT_FILE, data_only=True)

    @classmethod
    def tearDownClass(cls):
        ai.INPUT_DIR, ai.OUTPUT_DIR, ai.OUTPUT_FILE = cls.old
        cls.wb.close()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_exit_code(self):
        self.assertEqual(self.exit_code, 0)

    def test_output_has_all_required_sheets(self):
        for name in ("Integrated Master", "Integration_Log", "Unmatched_Students",
                     "Unmatched_Subjects", "Summary"):
            self.assertIn(name, self.wb.sheetnames)
        # only the semesters with data come first, Summary right after them
        expected = ["Semester 1", "Semester 2", "Summary"]
        self.assertEqual(self.wb.sheetnames[:3], expected)

    def test_download_keeps_only_the_semester_sheets(self):
        payload = ai.semester_only_bytes(ai.OUTPUT_FILE.read_bytes())
        wb = openpyxl.load_workbook(io.BytesIO(payload))
        self.assertEqual(wb.sheetnames, ["Semester 1", "Semester 2"])
        wb.close()

    def test_download_without_semester_sheets_is_unchanged(self):
        wb = openpyxl.Workbook()
        wb.active.title = ai.SUMMARY_SHEET
        buffer = io.BytesIO()
        wb.save(buffer)
        wb.close()
        payload = buffer.getvalue()
        self.assertEqual(ai.semester_only_bytes(payload), payload)

    # ---------------------------------------------------------- Semester sheets
    def _sem(self, sem):
        """(row 3 headers, student rows) of one semester sheet."""
        ws = self.wb[f"Semester {sem}"]
        headers = [c.value for c in ws[3]]
        rows = [[c.value for c in r]
                for r in ws.iter_rows(min_row=4, max_col=len(headers))]
        return headers, rows

    @staticmethod
    def _row(rows, usn):
        return next(r for r in rows if str(r[3]) == usn)

    @staticmethod
    def _subject(row, index):
        """Allotted, Total, A, B, A+B, % of the index-th subject section."""
        base = len(ai.SEM_INFO_COLUMNS) + index * len(ai.SEM_SUBJECT_COLUMNS)
        return row[base:base + len(ai.SEM_SUBJECT_COLUMNS)]

    @staticmethod
    def _count(value):
        return int(float(value))

    @staticmethod
    def _percent(value):
        return round(float(value), 2)

    def test_semester_sheet_has_required_columns(self):
        for sem in (1, 2):                       # the two semesters with data
            headers, _ = self._sem(sem)
            self.assertEqual(headers[:len(ai.SEM_INFO_COLUMNS)], ai.SEM_INFO_COLUMNS,
                             f"Semester {sem} info columns")
            rest = headers[len(ai.SEM_INFO_COLUMNS):]
            self.assertEqual(len(rest) % len(ai.SEM_SUBJECT_COLUMNS), 0,
                             f"Semester {sem} subject columns")
            for i in range(len(rest) // len(ai.SEM_SUBJECT_COLUMNS)):
                chunk = rest[i * 6:(i + 1) * 6]
                self.assertEqual(chunk, ai.SEM_SUBJECT_COLUMNS,
                                 f"Semester {sem} subject block {i + 1}")

    def test_semester1_follows_the_college_template(self):
        headers, rows = self._sem(1)
        ws = self.wb["Semester 1"]
        self.assertEqual(ws.cell(2, 8).value, "1st Semester")   # title in H2
        self.assertEqual(ws.cell(2, 11).value, "Sub: Mathematics BTE1001")
        self.assertEqual(ws.cell(2, 17).value, "Sub: BTE9999")
        self.assertEqual(len(headers), 10 + 2 * 6)

        alpha = self._row(rows, "AJU/1001")
        self.assertEqual(alpha[:7],
                         [1, "ALPHA STUDENT", None, "AJU/1001", None, None, "REGULAR"])
        # BTE1001 : 60 classes, 40 attended -> pushed into 75-78% (46/60 = 76.67%)
        self.assertEqual(self._subject(alpha, 0), [60, 60, 40, 6, 46, 76.67])
        # unknown course BTE9999 is still counted: 40/20 -> 31/40 = 77.5%
        self.assertEqual(self._subject(alpha, 1), [40, 40, 20, 11, 31, 77.5])
        # student level: weighted sums of the two subject sections
        self.assertEqual(self._count(alpha[7]), 100)            # Total No. of Classess
        self.assertEqual(self._count(alpha[8]), 77)             # Attended (A+B)
        self.assertEqual(self._percent(alpha[9]), 77.0)         # Current Attendance %

        # a student without usable data keeps the row but no invented numbers
        beta = self._row(rows, "AJU/1002")
        self.assertEqual(beta[:7],
                         [2, "BETA STUDENT", None, "AJU/1002", None, None, "REGULAR"])
        self.assertEqual(beta[7], None)
        self.assertEqual(beta[8], None)
        self.assertEqual(self._subject(beta, 0), [None] * 6)

    def test_semester_total_is_weighted_not_averaged(self):
        _, rows = self._sem(1)
        alpha = self._row(rows, "AJU/1001")
        subjects = [self._subject(alpha, i) for i in range(2)]
        total = sum(s[1] for s in subjects)                 # Total No. of Classess
        final = sum(s[4] for s in subjects)                 # Attended (A+B)
        self.assertEqual(self._count(alpha[7]), total)
        self.assertEqual(self._count(alpha[8]), final)
        self.assertEqual(self._percent(alpha[9]),
                         round(final / total * 100, 2))
        # never the average of the subject percentages (66.67% and 50%)
        average_of_actuals = (40 / 60 * 100 + 20 / 40 * 100) / 2
        self.assertNotEqual(self._percent(alpha[9]), round(average_of_actuals, 2))
        # actual attendance stays visible next to the adjusted attendance
        self.assertEqual(self._count(sum(s[2] for s in subjects)), 60)
        self.assertEqual(self._count(final), 77)

    def test_semester2_applies_the_95_band_rule(self):
        headers, rows = self._sem(2)
        self.assertEqual(headers[:10], ai.SEM_INFO_COLUMNS)
        alpha = self._row(rows, "AJU/1001")
        # 90% -> pushed into 95-98% with the smallest possible number of classes
        self.assertEqual(self._subject(alpha, 0), [100, 100, 90, 5, 95, 95])
        self.assertEqual(self._count(alpha[7]), 100)
        self.assertEqual(self._count(alpha[8]), 95)
        self.assertEqual(self._percent(alpha[9]), 95.0)         # never 115

    def test_empty_semester_gets_no_sheet(self):
        for sem in (3, 4, 5, 6, 7):
            self.assertNotIn(f"Semester {sem}", self.wb.sheetnames,
                             f"Semester {sem} has no data and must get no sheet")

    # ------------------------------------------------------------- Summary sheet
    def _summary_rows(self):
        ws = self.wb["Summary"]
        return [[c.value for c in r] for r in ws.iter_rows()]

    def test_summary_overall_table_is_weighted(self):
        rows = self._summary_rows()
        header_i = next(i for i, r in enumerate(rows) if r[0] == "Student Name")
        self.assertEqual(rows[header_i][:12], ai.OVERALL_COLUMNS)
        alpha = next(r for r in rows[header_i + 1:] if r[1] == "AJU/1001")
        self.assertEqual(alpha[4], "1,2")                 # semesters with data
        self.assertEqual(self._count(alpha[5]), 200)      # 100 + 100
        self.assertEqual(self._count(alpha[6]), 150)      # 60 + 90
        self.assertEqual(self._percent(alpha[7]), 75.0)   # 150/200, not an average
        self.assertEqual(self._count(alpha[8]), 17)       # 6 + 11 required classes
        self.assertEqual(self._percent(alpha[9]), 86.0)   # 172/200 adjusted
        self.assertEqual(self._count(alpha[10]), 172)
        self.assertEqual(self._count(alpha[11]), 5)       # semester 2 bonus classes
        beta = next(r for r in rows[header_i + 1:] if r[1] == "AJU/1002")
        self.assertIn(beta[5], ("", None))                # no data -> blank, not 0%

    def test_summary_semester_table(self):
        rows = self._summary_rows()
        header_i = next(i for i, r in enumerate(rows) if r[0] == "Semester")
        table = {r[0]: r for r in rows[header_i + 1:header_i + 4]}
        self.assertEqual(len(table), 3)           # 2 data semesters + ALL SEMESTERS
        sem1 = table["Semester 1"]
        self.assertEqual(self._count(sem1[1]), 1)     # students
        self.assertEqual(self._count(sem1[2]), 100)   # total classes
        self.assertEqual(self._count(sem1[3]), 60)    # attended
        self.assertEqual(self._percent(sem1[4]), 60.0)
        self.assertEqual(self._count(sem1[5]), 17)    # required classes
        self.assertEqual(self._count(sem1[6]), 0)     # bonus classes
        self.assertEqual(self._count(sem1[7]), 77)    # final attended
        self.assertEqual(self._percent(sem1[8]), 77.0)
        for sem in (3, 4, 5, 6, 7):
            self.assertNotIn(f"Semester {sem}", table,
                             "an empty semester must not get a summary row")
        overall = table["ALL SEMESTERS"]
        self.assertEqual(self._count(overall[2]), 200)
        self.assertEqual(self._count(overall[3]), 150)
        self.assertEqual(self._count(overall[5]), 17)
        self.assertEqual(self._count(overall[6]), 5)
        self.assertEqual(self._count(overall[7]), 172)
        self.assertEqual(self._percent(overall[8]), 86.0)

    def test_summary_reports_validation_pass(self):
        rows = self._summary_rows()
        verdict = next(r[1] for r in rows if r and r[0] == "Validation")
        self.assertTrue(str(verdict).startswith("PASS"), verdict)

    def test_master_sheet_is_preserved(self):
        ws = self.wb["Integrated Master"]
        self.assertEqual(ws.cell(4, 3).value, "AJU/1001")   # original data intact
        self.assertEqual(ws.cell(4, 2).value, "ALPHA STUDENT")

    def test_subject_block_written(self):
        ws = self.wb["Integrated Master"]
        layout = ai.parse_master_layout(ws)
        block = layout.subject_index[(1, "BTE1001")]
        row = layout.student_rows[ai.student_key("AJU/1001")]
        # the subject section carries the six columns of the college template
        headers = {role: ws.cell(layout.header_row, col).value
                   for role, col in block.columns.items()}
        self.assertEqual(headers["allotted"], "Total No of Allotted Classes")
        self.assertEqual(headers["total"], "Total No. of Classess")
        self.assertEqual(headers["actual"],
                         "Total No of Classes Attended- Actual (A)")
        self.assertEqual(headers["added"], "Total No of Classes Added (B)")
        self.assertEqual(headers["final"], "Total No of Classess Attended (A+B)")
        self.assertEqual(headers["pct"], "Current Attendance %")
        # BTE1001 : allotted / taken 60, attended 40, B = 6, A+B = 46, 76.67%
        self.assertEqual(ws.cell(row, block.columns["allotted"]).value, 60)
        self.assertEqual(ws.cell(row, block.columns["total"]).value, 60)
        self.assertEqual(ws.cell(row, block.columns["actual"]).value, 40)
        self.assertEqual(ws.cell(row, block.columns["added"]).value, 6)
        self.assertEqual(ws.cell(row, block.columns["final"]).value, 46)
        self.assertEqual(round(ws.cell(row, block.columns["pct"]).value, 2), 76.67)

    def test_master_semester_columns_follow_the_stated_meanings(self):
        """Every semester column of the MASTER holds exactly what it says."""
        ws = self.wb["Integrated Master"]
        layout = ai.parse_master_layout(ws)
        block = layout.semester_blocks[1]
        row = layout.student_rows[ai.student_key("AJU/1001")]
        headers = {role: ws.cell(layout.header_row, col).value
                   for role, col in block.columns.items()}
        self.assertEqual(headers["total"], "Total No. of Classess")
        self.assertEqual(headers["actual"],
                         "Total No of Actual  Clases Attended (A)")
        self.assertEqual(headers["added"], "Total No of Additional Classes(B)")
        self.assertEqual(headers["final"], "Total No. of Classess Attended(A+B)")
        self.assertEqual(headers["pct"], "Current Attendance %")

        # subject sections of the Semester 1 sheet are the source of the totals
        _, sem_rows = self._sem(1)
        alpha = self._row(sem_rows, "AJU/1001")
        subjects = [self._subject(alpha, i) for i in range(2)]
        self.assertEqual(len(subjects), 2)

        total = sum(s[1] for s in subjects)              # all subjects
        attended = sum(s[2] for s in subjects)           # actually attended (A)
        added = sum(s[3] for s in subjects)              # B added by the rules
        final = sum(s[4] for s in subjects)              # A+B
        self.assertEqual((total, attended, added, final), (100, 60, 17, 77))

        cell = lambda role: ws.cell(row, block.columns[role]).value
        self.assertEqual(cell("total"), total)          # Total No. of Classess
        self.assertEqual(cell("actual"), attended)      # Actual Clases Attended (A)
        self.assertEqual(cell("added"), added)          # Additional Classes (B)
        self.assertEqual(cell("final"), final)          # Attended (A+B)
        self.assertEqual(round(cell("pct"), 2),
                         round(final / total * 100, 2))     # Current Attendance %
        # the semester sheet carries the same A+B figure at student level
        self.assertEqual(self._count(alpha[8]), final)

    def test_semester_block_percent_is_the_weighted_percentage(self):
        ws = self.wb["Integrated Master"]
        layout = ai.parse_master_layout(ws)
        block = layout.semester_blocks[2]               # semester 2: 90% -> 95%
        row = layout.student_rows[ai.student_key("AJU/1001")]
        total = ws.cell(row, block.columns["total"]).value
        actual = ws.cell(row, block.columns["actual"]).value
        added = ws.cell(row, block.columns["added"]).value
        final = ws.cell(row, block.columns["final"]).value
        pct = ws.cell(row, block.columns["pct"]).value
        self.assertEqual((total, actual, added, final), (100, 90, 5, 95))
        self.assertEqual(final, actual + added)
        self.assertAlmostEqual(pct, final / total * 100)     # 95%, never 115%
        self.assertLessEqual(pct, 100)

    def test_semester_block_is_weighted(self):
        ws = self.wb["Integrated Master"]
        # semester 1 = 46/60 (band) + 31/40 (band) => 77/100 = 77%
        self.assertEqual(ws.cell(4, 6).value, 100)     # total classes
        self.assertEqual(ws.cell(4, 7).value, 60)      # actual
        self.assertEqual(ws.cell(4, 8).value, 17)      # 6 + 11 band classes
        self.assertEqual(ws.cell(4, 9).value, 77)      # A+B
        self.assertEqual(round(ws.cell(4, 10).value, 2), 77.0)

    def test_overall_block_is_weighted(self):
        ws = self.wb["Integrated Master"]
        layout = ai.parse_master_layout(ws)
        overall = layout.overall_block
        self.assertIsNotNone(overall)
        cols = overall.columns
        total = ws.cell(4, cols["total"]).value
        final = ws.cell(4, cols["final"]).value
        pct = ws.cell(4, cols["pct"]).value
        self.assertEqual(total, 200)                  # 100 semester 1 + 100 semester 2
        self.assertEqual(round(final / total * 100, 2), round(pct, 2))

    def test_unmatched_student_reported(self):
        ws = self.wb["Unmatched_Students"]
        usns = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
        self.assertIn("AJU/7777", usns)

    def test_unmatched_subject_reported(self):
        ws = self.wb["Unmatched_Subjects"]
        codes = {ws.cell(r, 3).value for r in range(2, ws.max_row + 1)}
        self.assertIn("BTE9999", codes)

    def test_log_covers_every_source_record(self):
        ws = self.wb["Integration_Log"]
        headers = [c.value for c in ws[1]]
        self.assertEqual(len(headers), len(ai.LOG_COLUMNS))
        statuses = {ws.cell(r, headers.index("Match Status") + 1).value
                    for r in range(2, ws.max_row + 1)}
        self.assertIn("MATCHED", statuses)
        self.assertIn("STUDENT_NOT_FOUND", statuses)
        self.assertIn("COURSE_NOT_IN_MASTER", statuses)
        self.assertIn("MISSING_TOTAL", statuses)

    def test_output_validates(self):
        self.assertEqual(ai.validate_output_file(ai.OUTPUT_FILE), [])


# --------------------------------------------------------------------------- #
# 13, 14: the bands in the workbook + the flow without a MASTER file
# --------------------------------------------------------------------------- #
class RosterModeTest(unittest.TestCase):
    """No MASTER in the input folder -> the roster comes from the exports."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="attendance_roster_"))
        (cls.tmp / "input").mkdir()
        (cls.tmp / "output").mkdir()
        make_attendance(cls.tmp / "input" / "Semester_1.xlsx")
        cls.old = (ai.INPUT_DIR, ai.OUTPUT_DIR, ai.OUTPUT_FILE)
        ai.INPUT_DIR = cls.tmp / "input"
        ai.OUTPUT_DIR = cls.tmp / "output"
        ai.OUTPUT_FILE = cls.tmp / "output" / "Integrated_Attendance.xlsx"
        cls.log = io.StringIO()
        with argv_is([str(ai.INPUT_DIR)]), contextlib.redirect_stdout(cls.log):
            cls.exit_code = ai.main()
        cls.wb = openpyxl.load_workbook(ai.OUTPUT_FILE, data_only=True)

    @classmethod
    def tearDownClass(cls):
        ai.INPUT_DIR, ai.OUTPUT_DIR, ai.OUTPUT_FILE = cls.old
        cls.wb.close()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_exit_code_and_validation_pass(self):
        self.assertEqual(self.exit_code, 0)
        self.assertEqual(ai.validate_output_file(ai.OUTPUT_FILE), [])
        self.assertIn("Validation                    : PASS", self.log.getvalue())

    def test_run_is_reported_as_roster_without_master(self):
        text = self.log.getvalue()
        self.assertIn("Master file          : none", text)
        self.assertIn("Students in roster (no MASTER)", text)

    def test_integrated_master_sheet_is_the_roster(self):
        ws = self.wb[ai.MASTER_SHEET_TITLE]
        headers = [ws.cell(2, c).value
                   for c in range(1, len(ai.ROSTER_INFO_HEADERS) + 1)]
        self.assertEqual(headers, ai.ROSTER_INFO_HEADERS)
        labels = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        self.assertIn("Semester -01", labels)
        self.assertIn("Semester -02", labels)
        usns = {ws.cell(r, 4).value for r in range(3, ws.max_row + 1)}
        self.assertEqual(usns, {"AJU/1001", "AJU/1002", "AJU/7777"})

    def test_semester1_shows_the_same_bands_as_the_master_flow(self):
        ws = self.wb["Semester 1"]
        base = len(ai.SEM_INFO_COLUMNS)
        for row in range(4, ws.max_row + 1):
            if str(ws.cell(row, 4).value) == "AJU/1001":
                subject = [ws.cell(row, base + 1 + i).value
                           for i in range(len(ai.SEM_SUBJECT_COLUMNS))]
                break
        else:                                       # pragma: no cover - defensive
            self.fail("AJU/1001 is missing from the roster-based Semester 1 sheet")
        # 40/60 (66.67%) -> 46/60 = 76.67%, inside the 75-78% band
        self.assertEqual(subject, [60, 60, 40, 6, 46, 76.67])

    def test_summary_names_the_generated_roster(self):
        ws = self.wb[ai.SUMMARY_SHEET]
        self.assertEqual(ws.cell(3, 1).value, "Master file")
        self.assertEqual(ws.cell(3, 2).value, "(built from the attendance files)")
        labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        self.assertIn("Attendance bands", labels)

    def test_everybody_in_the_roster_is_matched(self):
        ws = self.wb[ai.UNMATCHED_STUDENTS_SHEET]
        self.assertLessEqual(ws.max_row, 1)         # header only, nobody left out


class InputDiscoveryTest(unittest.TestCase):
    """Folder argument, otherwise the most recently updated attendance folder."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="attendance_input_"))
        self.old_input = ai.INPUT_DIR
        ai.INPUT_DIR = self.tmp / "input"
        ai.INPUT_DIR.mkdir(parents=True)

    def tearDown(self):
        ai.INPUT_DIR = self.old_input
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_folder_argument_is_used(self):
        target = self.tmp / "explicit"
        target.mkdir()
        self.assertEqual(ai.resolve_input_dir([str(target)]), target)

    def test_file_argument_uses_its_folder(self):
        target = self.tmp / "explicit"
        target.mkdir()
        file = target / "Attendance.xlsx"
        file.write_bytes(b"")
        self.assertEqual(ai.resolve_input_dir([str(file)]), target)

    def test_missing_argument_warns_and_falls_back(self):
        with chdir(self.tmp), contextlib.redirect_stdout(io.StringIO()) as out:
            resolved = ai.resolve_input_dir([str(self.tmp / "nope")])
        self.assertIn("folder not found", out.getvalue())
        self.assertEqual(resolved, ai.INPUT_DIR)    # empty -> no folders found

    def test_newest_attendance_folder_wins(self):
        older = self.tmp / "older"
        newer = self.tmp / "newer"
        for folder in (older, newer):
            folder.mkdir()
            make_attendance(folder / "Attendance.xlsx")
        stamp = 1_700_000_000
        os.utime(older / "Attendance.xlsx", (stamp, stamp))
        os.utime(newer / "Attendance.xlsx", (stamp + 100, stamp + 100))
        with chdir(self.tmp), contextlib.redirect_stdout(io.StringIO()):
            resolved = ai.resolve_input_dir([])
        self.assertEqual(resolved, newer.resolve())

    def test_files_are_found_in_sub_folders(self):
        nested = self.tmp / "AIML" / "Attendance"
        nested.mkdir(parents=True)
        make_attendance(nested / "Attendance_Sem-I.xlsx")
        (self.tmp / "readme.txt").write_text("ignored", encoding="utf-8")
        found = ai.load_excel_files(self.tmp)
        self.assertEqual([p.parent for p in found], [nested])


def make_batch_attendance(path: Path, usns) -> None:
    """One export row per student - used for the admission year filter."""
    headers = ["USN NO", "STUDENT TYPE", "STUDENT_NAME", "SEMESTER", "SECTION",
               "COURSE_CODE1", "ATTENDANCE_FOR1", "TOTAL_CLASS1",
               "ATTENDED_CLASS1", "PERCENTAGE1"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(headers)
    for i, usn in enumerate(usns, start=1):
        ws.append([usn, "REGULAR", f"STUDENT {i}", "I", "C",
                   "BTE1001", "THEORY", 60, 40, 67])
    wb.save(path)


class OtherBatchTest(unittest.TestCase):
    """Students whose USN carries another admission year leave the roster."""

    def test_usn_batch_year_reads_the_admission_year(self):
        self.assertEqual(ai.usn_batch_year("AJU/240916"), 2024)
        self.assertEqual(ai.usn_batch_year("aju/230885"), 2023)
        self.assertEqual(ai.usn_batch_year("AJU/252824"), 2025)
        self.assertIsNone(ai.usn_batch_year("AJU/1001"))        # too short
        self.assertIsNone(ai.usn_batch_year("1AY23CS045"))      # no <YY><serial> tail
        self.assertIsNone(ai.usn_batch_year(""))

    def _records(self, usns):
        tmp = Path(tempfile.mkdtemp(prefix="attendance_batch_"))
        self.addCleanup(shutil.rmtree, tmp, True)
        path = tmp / "Attendance.xlsx"
        make_batch_attendance(path, usns)
        return ai.read_attendance_workbook(path).records

    def test_majority_year_is_detected_and_minority_is_dropped(self):
        records = self._records(["AJU/250001", "AJU/250002", "AJU/250003",
                                 "AJU/240099"])
        kept, rows = ai.split_off_other_batch(records)
        self.assertEqual({r.usn for r in kept},
                         {"AJU/250001", "AJU/250002", "AJU/250003"})
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], "AJU/240099")
        self.assertEqual(rows[0][2], 2024)                    # USN batch year
        self.assertEqual(rows[0][3], 2025)                      # export batch year

    def test_unreadable_usns_are_kept(self):
        records = self._records(["AJU/250001", "AJU/250002", "AJU/7777"])
        kept, rows = ai.split_off_other_batch(records)
        self.assertEqual({r.usn for r in kept},
                         {"AJU/250001", "AJU/250002", "AJU/7777"})
        self.assertEqual(rows, [])

    def test_a_folder_that_mixes_batches_is_left_alone(self):
        records = self._records(["AJU/250001", "AJU/240001"])
        kept, rows = ai.split_off_other_batch(records)
        self.assertEqual(len(kept), 2)
        self.assertEqual(rows, [])

    def test_master_flow_never_filters(self):
        usns = ["AJU/250001", "AJU/250002", "AJU/240099"]
        records = self._records(usns)
        tmp = Path(tempfile.mkdtemp(prefix="attendance_master_"))
        self.addCleanup(shutil.rmtree, tmp, True)
        make_master(tmp / "Master.xlsx",
                    [(1, "STUDENT 1", usns[0]), (2, "STUDENT 2", usns[1]),
                     (3, "STUDENT 3", usns[2])])
        data = ai.AttendanceFileData(path=tmp / "Attendance.xlsx", records=records)
        result = ai.integrate_attendance(tmp / "Master.xlsx", [data])
        self.assertFalse(result.master_generated)
        self.assertEqual(result.other_batch_students, [])
        self.assertEqual(result.stats["students_other_batch"], 0)
        # the MASTER defines the roster, the older batch student stays in it
        self.assertEqual(result.stats["students_matched"], 3)


class OtherBatchEndToEndTest(unittest.TestCase):
    """Full roster run: the older batch student is reported, not written."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="attendance_batch_run_"))
        (cls.tmp / "input").mkdir()
        (cls.tmp / "output").mkdir()
        make_batch_attendance(
            cls.tmp / "input" / "Attendance.xlsx",
            ["AJU/250001", "AJU/250002", "AJU/250003", "AJU/240099"])
        cls.old = (ai.INPUT_DIR, ai.OUTPUT_DIR, ai.OUTPUT_FILE)
        ai.INPUT_DIR = cls.tmp / "input"
        ai.OUTPUT_DIR = cls.tmp / "output"
        ai.OUTPUT_FILE = cls.tmp / "output" / "Integrated_Attendance.xlsx"
        cls.log = io.StringIO()
        with argv_is([str(ai.INPUT_DIR)]), contextlib.redirect_stdout(cls.log):
            cls.exit_code = ai.main()
        cls.wb = openpyxl.load_workbook(ai.OUTPUT_FILE, data_only=True)

    @classmethod
    def tearDownClass(cls):
        ai.INPUT_DIR, ai.OUTPUT_DIR, ai.OUTPUT_FILE = cls.old
        cls.wb.close()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_output_still_validates(self):
        self.assertEqual(self.exit_code, 0)
        self.assertEqual(ai.validate_output_file(ai.OUTPUT_FILE), [])

    def test_roster_holds_the_batch_only(self):
        ws = self.wb[ai.MASTER_SHEET_TITLE]
        usns = [ws.cell(r, 4).value for r in range(3, ws.max_row + 1)]
        self.assertEqual(usns, ["AJU/250001", "AJU/250002", "AJU/250003"])

    def test_semester_sheet_holds_the_batch_only(self):
        ws = self.wb["Semester 1"]
        usns = [str(ws.cell(r, 4).value) for r in range(4, ws.max_row + 1)]
        self.assertEqual(usns, ["AJU/250001", "AJU/250002", "AJU/250003"])
        self.assertNotIn("AJU/240099", usns)

    def test_other_batch_sheet_reports_the_student(self):
        ws = self.wb[ai.OTHER_BATCH_SHEET]
        self.assertEqual([c.value for c in ws[1]], ai.OTHER_BATCH_HEADERS)
        self.assertEqual(ws.max_row, 2)                 # header + one student
        self.assertEqual(ws.cell(2, 1).value, "AJU/240099")
        self.assertEqual(ws.cell(2, 3).value, 2024)     # USN batch year
        self.assertEqual(ws.cell(2, 4).value, 2025)     # export batch year

    def test_summary_counts_the_excluded_student(self):
        ws = self.wb[ai.SUMMARY_SHEET]
        labels = {ws.cell(r, 1).value: ws.cell(r, 2).value
                  for r in range(1, ws.max_row + 1)}
        excluded = next(v for label, v in labels.items()
                        if label and label.startswith("Students excluded"))
        self.assertEqual(excluded, 1)
        self.assertEqual(labels["Total students (roster from the attendance files)"], 3)


if __name__ == "__main__":
    unittest.main(verbosity=2)
